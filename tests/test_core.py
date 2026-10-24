import io
import json
from pathlib import Path

import pytest
import fitz
from docx import Document

import docassist as core
import templates


@pytest.fixture
def isolated(tmp_path, monkeypatch):
    monkeypatch.setattr(core, "DATA", tmp_path)
    monkeypatch.setattr(core, "DB", tmp_path / "test.db")
    monkeypatch.setattr(core, "FILES", tmp_path / "uploads")
    core.init_db()
    return tmp_path


def test_versions_and_search(isolated):
    first = Document()
    first.add_paragraph("1.1. Срок проекта составляет десять дней.")
    out = io.BytesIO(); first.save(out)
    v1 = core.ingest(out.getvalue(), "regulation.docx")
    assert any(x["revision_id"] == v1 for x in core.search("Срок проекта"))
    second = Document()
    second.add_paragraph("1.1. Срок проекта составляет двадцать дней.")
    out = io.BytesIO(); second.save(out)
    doc_id = core.documents()[0]["id"]
    v2 = core.ingest(out.getvalue(), "regulation.docx", doc_id)
    found = core.search("Срок проекта")
    assert found and all(x["revision_id"] == v2 for x in found)
    assert "двадцать" in found[0]["text"]
    assert len(core.revisions(doc_id)) == 2
    core.activate(doc_id, v1)
    assert "десять" in core.search("Срок проекта")[0]["text"]


def test_table_rows_keep_headers(isolated):
    path = next(core.SEED.glob("Матрица*.docx"))
    rid = core.ingest(path.read_bytes(), path.name)
    with core.connect() as db:
        rows = db.execute("SELECT text FROM fragments WHERE revision_id=? AND kind='table'", (rid,)).fetchall()
    assert any("Департамент ДОУ" in x["text"] and "временных объединений" in x["text"] for x in rows)


def test_pdf_page_location(isolated):
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((50, 50), "1.1. Document registration every day")
    data = pdf.tobytes()
    revision = core.ingest(data, "instruction.pdf")
    with core.connect() as db:
        row = db.execute("SELECT locator,section FROM fragments WHERE revision_id=?", (revision,)).fetchone()
    assert row["locator"] == "стр. 1"
    assert row["section"] == "1.1"


def test_approval_conditions():
    ordinary = {x["department"] for x in core.approval_route("working_group", False)}
    it = {x["department"] for x in core.approval_route("working_group", True)}
    system = {x["department"] for x in core.approval_route("information_system")}
    assert "Департамент ДОУ" in ordinary
    assert "Кадровое подразделение" in ordinary
    assert "Департамент информационных технологий" not in ordinary
    assert "Департамент информационных технологий" in it
    assert "Департамент информационной безопасности" in system
    assert "Юридическое подразделение" in system


def test_route_answer_has_all_required_units(isolated):
    path = next(core.SEED.glob("Матрица*.docx"))
    core.ingest(path.read_bytes(), path.name)
    result = core.answer("Кто визирует рабочую группу для ИТ-проекта?")
    assert result["status"] == "answered"
    assert len(result["sources"]) == 1
    for name in ("Департамент ДОУ", "Кадровое подразделение", "Департамент информационных технологий"):
        assert name in result["answer"]
    ordinary = core.answer("Кто согласует рабочую группу без ИТ-компонента?")
    assert "Департамент информационных технологий" not in ordinary["answer"]
    assert core.answer("Кто согласует рабочую группу?")["status"] == "need_clarification"
    system = core.answer("Кто визирует приказ о вводе информационной системы?")
    for name in ("Юридическое подразделение", "Департамент ДОУ", "Департамент информационной безопасности", "Департамент информационных технологий"):
        assert name in system["answer"]


@pytest.mark.parametrize("item", templates.catalog())
def test_all_original_modules(item):
    values = {key: f"Значение {key}" for key in item["notes"]}
    values.update({"3_leader": "Руководитель", "3_secretary": "Секретарь"})
    data = templates.create_draft(item["id"], values)
    original = Document(templates.template_path(item["id"]))
    draft = Document(io.BytesIO(data))
    assert len(draft.tables) == len(original.tables)
    assert "Примечания по заполнению" not in "\n".join(p.text for p in draft.paragraphs)
    assert "ДАТА / РЕГИСТРАЦИОННЫЙ НОМЕР" in "\n".join(p.text for p in draft.paragraphs)
    assert "ВИЗУАЛИЗАЦИЯ ЭП" in "\n".join(p.text for p in draft.paragraphs) or any(
        "ВИЗУАЛИЗАЦИЯ ЭП" in c.text for t in draft.tables for r in t.rows for c in r.cells)


def test_missing_fields_rejected():
    item = templates.catalog()[0]
    with pytest.raises(ValueError):
        templates.create_draft(item["id"], {})


def test_provider_error_does_not_fake_answer(isolated):
    d = Document(); d.add_paragraph("1.1. Регистрация документов выполняется ежедневно.")
    out = io.BytesIO(); d.save(out)
    core.ingest(out.getvalue(), "rule.docx")
    class Offline:
        def complete(self, messages):
            raise core.GigaChatError("GigaChat недоступен")
    result = core.answer("Регистрация документов", client=Offline())
    assert result["status"] == "provider_error"
    assert not result["sources"]


def test_unknown_question(isolated):
    assert core.answer("квантовый виноград") ["status"] == "not_found"


def test_dialogue_keeps_revision_links(isolated):
    session = core.new_session()
    core.add_message(session, "user", "Первый вопрос")
    core.add_message(session, "assistant", "Ответ [1]", [7], [11])
    saved = core.messages(session)
    assert [x["role"] for x in saved] == ["user", "assistant"]
    assert json.loads(saved[-1]["revision_ids"]) == [7]
    assert json.loads(saved[-1]["source_ids"]) == [11]

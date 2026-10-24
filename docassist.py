"""Document store, retrieval and GigaChat integration for the demo application."""
from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import sqlite3
import time
import uuid
from datetime import datetime, timezone
from functools import lru_cache
from pathlib import Path

import fitz
import requests
from docx import Document

DATA = Path(os.getenv("DOCASSIST_DATA", "data"))
DB = DATA / "docassist.db"
FILES = DATA / "uploads"
SEED = DATA / "seed"
TOKEN_RE = re.compile(r"[\wа-яё]{3,}", re.I)


def connect():
    DATA.mkdir(parents=True, exist_ok=True)
    db = sqlite3.connect(DB)
    db.row_factory = sqlite3.Row
    db.execute("PRAGMA foreign_keys=ON")
    return db


def init_db():
    """Add new tables without changing or deleting the prototype's tables."""
    with connect() as db:
        db.executescript("""
        CREATE TABLE IF NOT EXISTS documents (
          id TEXT PRIMARY KEY, title TEXT NOT NULL, active_revision INTEGER);
        CREATE TABLE IF NOT EXISTS revisions (
          id INTEGER PRIMARY KEY, document_id TEXT NOT NULL REFERENCES documents(id),
          number INTEGER NOT NULL, filename TEXT NOT NULL, path TEXT NOT NULL,
          sha256 TEXT NOT NULL, indexed_at TEXT NOT NULL, status TEXT NOT NULL,
          UNIQUE(document_id, number));
        CREATE TABLE IF NOT EXISTS fragments (
          id INTEGER PRIMARY KEY, revision_id INTEGER NOT NULL REFERENCES revisions(id),
          section TEXT, locator TEXT, kind TEXT NOT NULL, text TEXT NOT NULL);
        CREATE VIRTUAL TABLE IF NOT EXISTS fragments_fts USING fts5(
          text, content='fragments', content_rowid='id', tokenize='unicode61');
        CREATE TABLE IF NOT EXISTS sessions (id TEXT PRIMARY KEY, created_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS messages (
          id INTEGER PRIMARY KEY, session_id TEXT NOT NULL REFERENCES sessions(id),
          role TEXT NOT NULL, content TEXT NOT NULL, revision_ids TEXT NOT NULL DEFAULT '[]',
          source_ids TEXT NOT NULL DEFAULT '[]',
          created_at TEXT NOT NULL);
        """)


def parse_file(path: Path):
    """Yield (section, locator, kind, text), keeping DOCX table rows together."""
    if path.suffix.lower() == ".pdf":
        with fitz.open(path) as pdf:
            for number, page in enumerate(pdf, 1):
                lines = [x.strip() for x in page.get_text().splitlines() if x.strip()]
                section, block = "", []
                for line in lines:
                    match = re.match(r"^(\d+(?:\.\d+)*\.?)\s+", line)
                    if match and block:
                        yield section, f"стр. {number}", "pdf", " ".join(block)
                        block = []
                    if match:
                        section = match.group(1).rstrip(".")
                    block.append(line)
                    if sum(map(len, block)) > 900:
                        yield section, f"стр. {number}", "pdf", " ".join(block)
                        block = []
                if block:
                    yield section, f"стр. {number}", "pdf", " ".join(block)
    elif path.suffix.lower() == ".docx":
        doc = Document(path)
        for i, paragraph in enumerate(doc.paragraphs):
            value = paragraph.text.strip()
            if value:
                match = re.match(r"^(\d+(?:\.\d+)*\.?)\s+", value)
                yield match.group(1).rstrip(".") if match else "", f"абзац {i + 1}", "docx", value
        for ti, table in enumerate(doc.tables, 1):
            headers = [c.text.strip() for c in table.rows[0].cells] if table.rows else []
            for ri, row in enumerate(table.rows[1:], 2):
                cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                if any(cells):
                    pairs = [f"{h}: {v}" for h, v in zip(headers, cells) if v]
                    yield "таблица", f"таблица {ti}, строка {ri}", "table", " | ".join(pairs)
    else:
        raise ValueError("Поддерживаются PDF и DOCX")


def ingest(content: bytes, filename: str, document_id: str | None = None):
    if not filename.lower().endswith((".pdf", ".docx")):
        raise ValueError("Поддерживаются PDF и DOCX")
    if len(content) > 25_000_000:
        raise ValueError("Файл больше 25 МБ")
    digest = hashlib.sha256(content).hexdigest()
    doc_id = document_id or hashlib.sha256(filename.casefold().encode()).hexdigest()[:20]
    fragments = list(parse_file_from_bytes(content, filename))
    if not fragments:
        raise ValueError("Из документа не извлечён текст; активная редакция сохранена")
    FILES.mkdir(parents=True, exist_ok=True)
    safe_name = re.sub(r"[^\w.\-]", "_", Path(filename).name)
    with connect() as db:
        db.execute("INSERT OR IGNORE INTO documents(id,title) VALUES(?,?)", (doc_id, filename))
        old = db.execute("SELECT COALESCE(MAX(number),0) FROM revisions WHERE document_id=?", (doc_id,)).fetchone()[0]
        duplicate = db.execute("SELECT id FROM revisions WHERE document_id=? AND sha256=?", (doc_id, digest)).fetchone()
        if duplicate:
            return duplicate[0]
        number = old + 1
        target = FILES / f"{doc_id}-{number}-{safe_name}"
        target.write_bytes(content)
        try:
            cur = db.execute("INSERT INTO revisions(document_id,number,filename,path,sha256,indexed_at,status) VALUES(?,?,?,?,?,?,?)",
                (doc_id, number, filename, str(target), digest, datetime.now(timezone.utc).isoformat(), "ready"))
            revision_id = cur.lastrowid
            for section, locator, kind, value in fragments:
                row = db.execute("INSERT INTO fragments(revision_id,section,locator,kind,text) VALUES(?,?,?,?,?)",
                    (revision_id, section, locator, kind, value))
                db.execute("INSERT INTO fragments_fts(rowid,text) VALUES(?,?)", (row.lastrowid, value))
            db.execute("UPDATE documents SET active_revision=? WHERE id=?", (revision_id, doc_id))
        except Exception:
            target.unlink(missing_ok=True)
            raise
    return revision_id


def parse_file_from_bytes(content: bytes, filename: str):
    # PyMuPDF and python-docx both accept binary streams; the temporary file
    # preserves one parser path for uploaded files and seed files.
    import tempfile
    with tempfile.TemporaryDirectory() as folder:
        path = Path(folder) / ("source" + Path(filename).suffix.lower())
        path.write_bytes(content)
        yield from parse_file(path)


def seed_documents():
    init_db()
    for path in SEED.iterdir() if SEED.exists() else []:
        if path.suffix.lower() in (".pdf", ".docx"):
            doc_id = hashlib.sha256(path.name.casefold().encode()).hexdigest()[:20]
            with connect() as db:
                exists = db.execute("SELECT 1 FROM documents WHERE id=?", (doc_id,)).fetchone()
            if not exists:
                ingest(path.read_bytes(), path.name, doc_id)


def documents():
    with connect() as db:
        return [dict(r) for r in db.execute("""SELECT d.id,d.title,d.active_revision,r.number,r.status,r.indexed_at
            FROM documents d LEFT JOIN revisions r ON r.id=d.active_revision ORDER BY d.title""")]


def revisions(document_id):
    with connect() as db:
        return [dict(r) for r in db.execute("SELECT * FROM revisions WHERE document_id=? ORDER BY number DESC", (document_id,))]


def activate(document_id, revision_id):
    with connect() as db:
        valid = db.execute("SELECT 1 FROM revisions WHERE id=? AND document_id=? AND status='ready'", (revision_id, document_id)).fetchone()
        if not valid:
            raise ValueError("Редакция не найдена или индекс не готов")
        db.execute("UPDATE documents SET active_revision=? WHERE id=?", (revision_id, document_id))


def search(query: str, limit=6):
    words = TOKEN_RE.findall(query.casefold())
    stop = {"как", "что", "кто", "где", "для", "при", "или", "это", "мне", "нужно", "какие", "какой", "создай"}
    words = [w for w in words if w not in stop][:12]
    if not words:
        return []
    endings = ("иями", "ями", "ого", "ему", "ыми", "ами", "иях", "иях", "ать", "ить", "уют", "ает", "ется", "ение", "ений", "ую", "ой", "ых", "ие", "ая", "ый", "ые", "ам", "ах", "ов", "ом", "у", "а", "ы", "и")
    stems = []
    for word in words:
        stem = word
        for ending in endings:
            if word.endswith(ending) and len(word) - len(ending) >= 4:
                stem = word[:-len(ending)]
                break
        stems.append(stem)
    if any(w.startswith("соглас") for w in words):
        stems.append("визир")
    fts_query = " OR ".join('"' + w.replace('"', '') + '"*' for w in stems)
    with connect() as db:
        rows = db.execute("""SELECT f.id,f.revision_id,f.section,f.locator,f.kind,f.text,d.title,
                   bm25(fragments_fts) AS score FROM fragments_fts
                   JOIN fragments f ON f.id=fragments_fts.rowid
                   JOIN revisions r ON r.id=f.revision_id
                   JOIN documents d ON d.active_revision=r.id
                   WHERE fragments_fts MATCH ? ORDER BY score LIMIT ?""", (fts_query, limit * 4)).fetchall()
    found = []
    for row in rows:
        item = dict(row)
        overlap = sum(w in item["text"].casefold() for w in stems)
        if overlap:
            item["score"] = overlap / len(words)
            found.append(item)
    found.sort(key=lambda r: r["score"], reverse=True)
    return found[:limit]


class GigaChatError(RuntimeError):
    pass


class GigaChat:
    def __init__(self, key=None, scope=None, model=None, session=None):
        self.key = key or os.getenv("GIGACHAT_AUTH_KEY", "")
        self.scope = scope or os.getenv("GIGACHAT_SCOPE", "GIGACHAT_API_PERS")
        self.model = model or os.getenv("GIGACHAT_MODEL", "GigaChat-3-Lightning")
        self.session = session or requests.Session()
        self.token = None
        self.expires = 0

    def _request(self, method, url, **kwargs):
        for attempt in range(3):
            try:
                response = self.session.request(method, url, timeout=(5, 35), **kwargs)
                if response.status_code in (429, 500, 502, 503, 504) and attempt < 2:
                    time.sleep(0.5 * 2 ** attempt)
                    continue
                response.raise_for_status()
                return response.json()
            except (requests.Timeout, requests.ConnectionError) as exc:
                if attempt == 2:
                    raise GigaChatError("GigaChat временно недоступен") from exc
                time.sleep(0.5 * 2 ** attempt)
            except (requests.HTTPError, ValueError) as exc:
                code = getattr(getattr(exc, "response", None), "status_code", None)
                raise GigaChatError(f"Ошибка GigaChat API ({code or 'неверный ответ'})") from exc
        raise GigaChatError("GigaChat временно недоступен")

    def access_token(self):
        if not self.key:
            raise GigaChatError("Ключ GigaChat не настроен")
        if self.token and time.time() < self.expires - 60:
            return self.token
        data = self._request("POST", "https://ngw.devices.sberbank.ru:9443/api/v2/oauth",
            headers={"Authorization": "Basic " + self.key, "RqUID": str(uuid.uuid4()), "Accept": "application/json"},
            data={"scope": self.scope})
        try:
            self.token = data["access_token"]
            self.expires = float(data.get("expires_at", time.time() + 1800))
        except (KeyError, TypeError, ValueError) as exc:
            raise GigaChatError("GigaChat вернул неверный токен") from exc
        return self.token

    def complete(self, messages):
        data = self._request("POST", "https://api.giga.chat/v1/chat/completions",
            headers={"Authorization": "Bearer " + self.access_token(), "Content-Type": "application/json"},
            json={"model": self.model, "messages": messages, "temperature": 0.1, "max_tokens": 1000})
        try:
            return data["choices"][0]["message"]["content"].strip()
        except (KeyError, IndexError, TypeError, AttributeError) as exc:
            raise GigaChatError("GigaChat вернул неверный ответ") from exc


@lru_cache(maxsize=1)
def default_client():
    return GigaChat()


def answer(question, history=(), client=None):
    context_question = " ".join([m["content"] for m in history[-3:] if m["role"] == "user"] + [question])
    lowered = context_question.casefold()
    if any(term in lowered for term in ("визир", "соглас", "маршрут", "кто")):
        route_kind = ("working_group" if "рабоч" in lowered and "групп" in lowered else
                      "information_system" if "ввод" in lowered and "информац" in lowered and "систем" in lowered else None)
        if route_kind:
            is_it = bool(re.search(r"\bит\b|ит-проект|информацион", question.casefold())) and not bool(
                re.search(r"\bне\s+(?:для\s+)?ит\b|без\s+ит\b|не\s+связан", question.casefold()))
            if route_kind == "working_group" and not is_it and not re.search(r"\bне\s+ит\b|без\s+ит\b|не\s+связан", lowered):
                return {"status": "need_clarification", "answer": "Рабочая группа создаётся для ИТ-проекта или решения вопросов информационных систем?", "sources": []}
            route = approval_route(route_kind, is_it)
            source = approval_source(route_kind)
            if route and source:
                lines = [f"- {x['department']} — {x['reason']}" for x in route]
                return {"status": "answered", "answer": "Согласно активной редакции матрицы обязательного визирования:\n" +
                        "\n".join(lines) + f"\n\nДополнительные подразделения могут потребоваться по затронутым компетенциям. [{source['id']}]",
                        "sources": [source]}
    hits = search(context_question)
    if not hits:
        return {"status": "not_found", "answer": "В загруженных документах не найдено подтверждение. Уточните вопрос или загрузите источник.", "sources": []}
    references = "\n".join(f"[{h['id']}] {h['title']}, {h['locator']}, пункт {h['section']}: {h['text'][:650]}" for h in hits)
    prompt = ("Ответь по-русски только по приведённым фрагментам. Если они не дают ответа, прямо скажи об этом. "
              "Не придумывай требования, пункты и согласующих. После каждого существенного вывода ставь ID источника в квадратных скобках.\n"
              f"Источники:\n{references[:6000]}\nВопрос: {question}")
    try:
        content = (client or default_client()).complete([{"role": "system", "content": "Ты помощник по внутренним документам. Не выходи за пределы источников."},
            {"role": "user", "content": prompt}])
    except GigaChatError as exc:
        return {"status": "provider_error", "answer": str(exc), "sources": []}
    cited = {int(x) for x in re.findall(r"\[(\d+)\]", content)}
    verified = [h for h in hits if h["id"] in cited]
    if not verified:
        return {"status": "not_found", "answer": "Не удалось подтвердить ответ ссылкой на найденный источник. Уточните вопрос.", "sources": []}
    return {"status": "answered", "answer": content, "sources": verified}


def new_session():
    session_id = uuid.uuid4().hex
    with connect() as db:
        db.execute("INSERT INTO sessions VALUES(?,?)", (session_id, datetime.now(timezone.utc).isoformat()))
    return session_id


def add_message(session_id, role, content, revision_ids=(), source_ids=()):
    with connect() as db:
        db.execute("INSERT INTO messages(session_id,role,content,revision_ids,source_ids,created_at) VALUES(?,?,?,?,?,?)",
            (session_id, role, content, json.dumps(list(revision_ids)), json.dumps(list(source_ids)), datetime.now(timezone.utc).isoformat()))


def messages(session_id):
    with connect() as db:
        return [dict(r) for r in db.execute("SELECT * FROM messages WHERE session_id=? ORDER BY id", (session_id,))]


def sessions():
    with connect() as db:
        return [dict(r) for r in db.execute("SELECT id,created_at FROM sessions ORDER BY created_at DESC LIMIT 30")]


def approval_route(kind, is_it=False):
    """Read a selected row of the source matrix; keep conditions explicit."""
    init_db()
    with connect() as db:
        active = db.execute("""SELECT r.path FROM documents d JOIN revisions r ON r.id=d.active_revision
            WHERE d.title LIKE 'Матрица%' ORDER BY d.title LIMIT 1""").fetchone()
    matrix = Path(active[0]) if active else next((p for p in SEED.glob("Матрица*.docx")), None)
    if not matrix:
        return []
    table = Document(matrix).tables[0]
    phrase = "вопросам информационных систем" if kind == "information_system" else "создании рабочих групп" if kind == "working_group" else None
    row_number = next((i for i, row in enumerate(table.rows) if phrase and phrase in row.cells[0].text.casefold()), None)
    if row_number is None:
        return []
    row = table.rows[row_number]
    labels = [c.text.strip() for c in table.rows[0].cells]
    answer = []
    for column in (1, 2, 3):
        if row.cells[column].text.strip().casefold() == "да":
            answer.append({"department": labels[column], "reason": f"Матрица, строка {row_number + 1}"})
    if kind == "information_system":
        for name in ("Департамент информационной безопасности", "Департамент информационных технологий"):
            if name in row.cells[4].text:
                answer.append({"department": name, "reason": f"Матрица, строка {row_number + 1}"})
    elif is_it and "Департамент информационных технологий" in row.cells[4].text:
        answer.append({"department": "Департамент информационных технологий",
                       "reason": f"Условие ИТ-проекта: матрица, строка {row_number + 1}"})
    return answer


def approval_source(kind):
    phrase = "создании рабочих групп" if kind == "working_group" else "вопросам информационных систем"
    with connect() as db:
        row = db.execute("""SELECT f.id,f.revision_id,f.section,f.locator,f.kind,f.text,d.title
            FROM fragments f JOIN revisions r ON r.id=f.revision_id
            JOIN documents d ON d.active_revision=r.id
            WHERE d.title LIKE 'Матрица%' AND f.kind='table' AND LOWER(f.text) LIKE ? LIMIT 1""",
            ("%" + phrase + "%",)).fetchone()
    return dict(row) if row else None
